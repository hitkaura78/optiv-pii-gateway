"""Seeded, labelled synthetic benchmark. Owner: builder.

    python -m pii_pipeline.benchmark [--out tmp/benchmark] [--seed 1] [--per-kind 9] [--kinds ...]

Writes a corpus (DOCX, PPTX, native PDF, scanned PDF rendered to an image, mixed PDF) plus
manifest.json into a gitignored folder, runs the real pipeline on every file and prints a per-type
precision/recall table, how the look-alike decoys were treated, and released/blocked/error counts
per format. Gold labels come from the generator, so they are exact. Only reserved/test values are
used (example.com/.org, 555-01xx numbers, test card numbers, 203.0.113.x / 2001:db8:: addresses).
SSNs have no reserved range with a valid area number, so widely published specimen/obviously fake
sequences are used. The manifest holds these synthetic values, so keep it out of Git (tmp/ is ignored).
"""
import argparse
import io
import json
import random
import sys
import time
from collections import Counter, defaultdict
from pathlib import Path

from . import cli, detect, evaluate, extract, ingest
from .models import Action, Finding, PIIType

KINDS = ("docx", "pptx", "pdf-native", "pdf-scanned", "pdf-mixed")
SUPPORTED = (PIIType.EMAIL, PIIType.PHONE, PIIType.CREDIT_CARD, PIIType.US_SSN, PIIType.IP_ADDRESS)
FIRST = ["alex", "sam", "kim", "lee", "maya", "noah", "priya", "omar"]
LAST = ["rivera", "chen", "okafor", "novak", "silva", "haddad", "kowalski"]
CARDS = ["4111 1111 1111 1111", "4242 4242 4242 4242", "5555 5555 5555 4444", "3782 822463 10005",
         "6011 1111 1111 1117", "4111-1111-1111-1111", "4111111111111111", "5105 1051 0510 5100"]
SSNS = ["123-45-6789", "321-54-9876", "219-09-9999", "078-05-1120", "111-22-3333", "234-56-7890"]
NAMES = ["Alex Rivera", "Maya Okafor", "Noah Novak", "Kim Silva", "Sam Chen", "Priya Haddad", "Omar Kowalski"]
STREETS = ["Evergreen Terrace", "Maple Avenue", "Harbor Road", "Baker Street", "Elm Court"]
CITIES = ["Springfield, IL 62704", "Portland, OR 97201", "London NW1 6XE", "Austin, TX 73301"]


def _email(r):
    f, l, n = r.choice(FIRST), r.choice(LAST), r.randint(10, 99)
    return r.choice([f"{f}.{l}{n}@example.com", f"{f}_{l}+{n}@example.org", f"{f[0]}{l}{n}@mail.example.com",
                     f"{f}.{l}{n}@EXAMPLE.ORG"])


def _phone(r, easy=False):
    n = f"01{r.randint(0, 99):02d}"
    strong = [f"415-555-{n}", f"(212) 555-{n}", f"+1 650 555 {n}", f"310.555.{n}", f"+44 20 7946 0{r.randint(100, 999)}"]
    return r.choice(strong if easy else strong + [f"555-{n}", f"415555{n}"])   # last two: weak by design


def _card(r):
    return r.choice(CARDS)


def _ssn(r):
    s = r.choice(SSNS)
    return r.choice([s, s.replace("-", "")])      # undashed 9 digits only counts with "SSN" context in the line


def _ip(r):
    return r.choice([f"203.0.113.{r.randint(1, 254)}", f"198.51.100.{r.randint(1, 254)}",
                     f"192.0.2.{r.randint(1, 254)}", f"2001:db8::{r.randint(1, 4095):x}",
                     "2001:db8:85a3::8a2e:370:7334"])


TRUE_GEN = {PIIType.EMAIL: _email, PIIType.PHONE: _phone, PIIType.CREDIT_CARD: _card,
            PIIType.US_SSN: _ssn, PIIType.IP_ADDRESS: _ip}
TRUE_LINES = {
    PIIType.EMAIL: ["Contact: {v}", "Please write to {v} for details.", "{v}"],
    PIIType.PHONE: ["Phone: {v}", "Call us on {v} after 9am.", "Mobile {v}"],
    PIIType.CREDIT_CARD: ["Card number {v}", "Paid with {v} (test card).", "{v}"],
    PIIType.US_SSN: ["SSN: {v}", "Social security number {v}", "SSN {v}"],
    PIIType.IP_ADDRESS: ["Server {v}", "Logged from {v} yesterday.", "{v}"],
}


EASY_DECOYS = {"short_number", "date_iso", "date_slash", "date_dot", "version", "amount", "ip_invalid"}


def _decoys(r, easy=False):
    y = r.randint(2019, 2026)
    return [
        ("order_number", f"ORD-{y}-{r.randint(100000, 999999)}", "Order {v} shipped."),
        ("short_number", f"#{r.randint(1000, 99999)}", "Ticket {v} closed."),
        ("date_iso", f"{y}-{r.randint(1, 12):02d}-{r.randint(1, 28):02d}", "Meeting on {v}."),
        ("date_slash", f"{r.randint(10, 28)}/{r.randint(10, 12)}/{y}", "Due {v}."),
        ("date_dot", f"{r.randint(10, 12)}.{r.randint(10, 28)}.{y}", "Updated {v}."),
        ("version", f"v{r.randint(1, 9)}.{r.randint(10, 99)}.0.{r.randint(1000, 9999)}", "Release {v} is out."),
        ("version_ip_like", f"{r.randint(1, 9)}.{r.randint(1, 9)}.{r.randint(1, 9)}.{r.randint(1, 9)}", "Build {v} passed."),
        ("isbn", f"978-{r.randint(0, 9)}-{r.randint(10, 99)}-{r.randint(100000, 999999)}-{r.randint(0, 9)}", "ISBN {v}"),
        ("luhn_invalid_card", r.choice(["4111 1111 1111 1112", "5500 0000 0000 0001", "4242 4242 4242 4243"]), "Card typo {v}"),
        ("ssn_invalid_area", r.choice(["000-12-3456", "666-12-3456", "912-34-5678"]), "Form id {v}"),
        ("ip_invalid", "999.1.1.1" if easy else r.choice(["999.1.1.1", "256.256.256.256", "300.10.10.10"]),
         "Bad address {v}"),
        ("amount", f"${r.randint(1, 9)},{r.randint(100, 999)},{r.randint(100, 999)}.{r.randint(10, 99)}", "Total {v}"),
    ]


def _unsupported(r):
    return [
        (PIIType.PERSON, r.choice(NAMES), "Signed by {v}"),
        (PIIType.ADDRESS, f"{r.randint(10, 999)} {r.choice(STREETS)}, {r.choice(CITIES)}", "Ship to {v}"),
        (PIIType.DATE_OF_BIRTH, f"{r.randint(1950, 2005)}-{r.randint(1, 12):02d}-{r.randint(1, 28):02d}",
         "Date of birth: {v}"),
    ]


def _clash(value, used):
    return any(value in u or u in value for u in used)


def make_items(rng: random.Random, easy: bool = False) -> dict:
    """Lines for one document: gold items (type, value, line), decoys (category, value, line).
    easy=True uses only well-formed phone numbers and decoys that should be silently ignored, so the
    corpus contains documents that can actually be released; the other documents are the hard cases."""
    used: list[str] = []
    gold, decoys = [], []

    def add(bucket, key, value, template):
        used.append(value)
        bucket.append({"key": key, "value": value, "line": template.format(v=value)})

    for t in rng.sample(SUPPORTED, rng.randint(4, 5)):
        for _ in range(rng.choice([1, 1, 2])):
            for _try in range(30):
                v = _phone(rng, easy) if t is PIIType.PHONE else TRUE_GEN[t](rng)
                if not _clash(v, used):
                    break
            else:
                continue
            tpl = rng.choice(TRUE_LINES[t])
            if t is PIIType.US_SSN and "-" not in v:
                tpl = "SSN: {v}"
            add(gold, t.value, v, tpl)
    pool = [d for d in _decoys(rng, easy) if not easy or d[0] in EASY_DECOYS]
    for cat, v, tpl in rng.sample(pool, min(len(pool), rng.randint(4, 6))):
        if not _clash(v, used):
            add(decoys, cat, v, tpl)
    for t, v, tpl in _unsupported(rng):
        if not _clash(v, used):
            add(gold, t.value, v, tpl)
    lines = [(g["line"]) for g in gold] + [d["line"] for d in decoys]
    rng.shuffle(lines)
    return {"gold": gold, "decoys": decoys, "lines": lines}


# ---- writers ------------------------------------------------------------------------------------

def _write_docx(path, lines):
    from docx import Document
    d = Document()
    for ln in lines:
        d.add_paragraph(ln)
    d.save(path)


def _write_pptx(path, lines):
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    for i in range(0, len(lines), 8):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        for k, ln in enumerate(lines[i:i + 8]):
            s.shapes.add_textbox(Inches(0.5), Inches(0.4 + 0.8 * k), Inches(9), Inches(0.6)).text_frame.text = ln
    prs.save(path)


def _png(lines):
    from PIL import Image, ImageDraw, ImageFont
    img = Image.new("RGB", (1654, 2339), "white")          # A4 at 200 dpi
    dr = ImageDraw.Draw(img)
    font = ImageFont.load_default(size=40)
    for k, ln in enumerate(lines):
        dr.text((120, 160 + 110 * k), ln, fill="black", font=font)
    buf = io.BytesIO()
    img.save(buf, "PNG")
    return buf.getvalue()


def _write_pdf(path, lines, mode):
    import pymupdf
    d = pymupdf.open()
    native, scanned = {"pdf-native": (lines, []), "pdf-scanned": ([], lines),
                       "pdf-mixed": (lines[:len(lines) // 2], lines[len(lines) // 2:])}[mode]
    if native or not scanned:
        p = d.new_page()
        for k, ln in enumerate(native):
            p.insert_text((60, 70 + 26 * k), ln, fontsize=12)
    if scanned:
        p = d.new_page()
        p.insert_image(p.rect, stream=_png(scanned))
    d.save(str(path))


def generate(out: Path, seed: int = 1, per_kind: int = 9, kinds=KINDS) -> dict:
    out = Path(out)
    corpus = out / "corpus"
    corpus.mkdir(parents=True, exist_ok=True)
    manifest = {"seed": seed, "files": []}
    for kind in kinds:
        for i in range(per_kind):
            rng = random.Random(f"{seed}/{kind}/{i}")
            items = make_items(rng, easy=(i % 3 == 0))
            if kind in ("pdf-scanned", "pdf-mixed"):
                items["lines"] = items["lines"][:14]               # keep one image page legible
                keep = " ".join(items["lines"])
                items["gold"] = [g for g in items["gold"] if g["line"] in items["lines"]]
                items["decoys"] = [d for d in items["decoys"] if d["line"] in keep]
            name = f"{kind}_{i:02d}.{'pdf' if kind.startswith('pdf') else kind}"
            path = corpus / name
            if kind == "docx":
                _write_docx(path, items["lines"])
            elif kind == "pptx":
                _write_pptx(path, items["lines"])
            else:
                _write_pdf(path, items["lines"], kind)
            manifest["files"].append({"file": name, "kind": kind, "gold": items["gold"], "decoys": items["decoys"]})
    (out / "manifest.json").write_text(json.dumps(manifest, indent=1), encoding="utf-8")
    return manifest


# ---- evaluation -----------------------------------------------------------------------------------

def _locate(doc, value, prefix, pii_type):
    """Find the unit holding a generator-known value; (Finding-like tuple) or None if extraction lost it."""
    for u in doc.units:
        i = u.text.find(value)
        if i >= 0:
            return prefix + u.unit_id, i, i + len(value)
    return None


def evaluate_file(entry, path: Path, work: Path):
    """Returns dict with predicted/gold Findings, decoy outcomes and the pipeline outcome."""
    res = {"kind": entry["kind"], "file": entry["file"], "predicted": [], "gold": [], "decoys": [],
           "outcome": None, "leaks": []}
    prefix = entry["file"] + "::"
    try:
        doc = extract.extract(ingest.load(path))
    except ValueError as exc:
        res["outcome"] = "error"
        res["error"] = str(exc)[:80]
        return res
    findings = detect.detect(doc)
    pred = [Finding(f.pii_type, prefix + f.unit_id, f.start, f.end, f.detector, f.confidence, f.action)
            for f in findings]
    res["predicted"] = pred
    for g in entry["gold"]:
        t = PIIType(g["key"])
        loc = _locate(doc, g["value"], prefix, t)
        # a value the extractor lost (e.g. OCR misread) stays gold with a unit that nothing matches: a miss
        unit, s, e = loc if loc else (prefix + "<not extracted>", 0, len(g["value"]))
        res["gold"].append(Finding(t, unit, s, e, "gold", 1.0, Action.KEEP))
    for d in entry["decoys"]:
        loc = _locate(doc, d["value"], prefix, None)
        if loc is None:
            res["decoys"].append((d["key"], "unlocated"))
            continue
        unit, s, e = loc
        hits = [p for p in pred if p.unit_id == unit and p.start is not None and p.start < e and s < p.end]
        kind = ("masked" if any(p.action in (Action.MASK, Action.REDACT) for p in hits)
                else "review" if hits else "clean")
        res["decoys"].append((d["key"], kind))
    try:
        decision, result, _, _ = cli.run(path, work, False)
    except (ValueError, OSError):
        res["outcome"] = "error"
        return res
    res["outcome"] = "released" if decision.passed else "blocked"
    if decision.passed:
        text = " ".join(u.text for u in extract.extract(ingest.load(result.output_path)).units)
        res["leaks"] = [g["key"] for g in entry["gold"] if g["value"] in text]
    return res


def run_corpus(out: Path) -> dict:
    out = Path(out)
    manifest = json.loads((out / "manifest.json").read_text(encoding="utf-8"))
    work = out / "work"
    work.mkdir(parents=True, exist_ok=True)
    predicted, gold = [], []
    outcomes: dict[str, Counter] = defaultdict(Counter)
    decoys: dict[str, Counter] = defaultdict(Counter)
    leaks: Counter = Counter()
    for entry in manifest["files"]:
        r = evaluate_file(entry, out / "corpus" / entry["file"], work)
        outcomes[entry["kind"]][r["outcome"]] += 1
        predicted += r["predicted"]
        if r["outcome"] != "error" or r["predicted"]:
            gold += r["gold"]
        for cat, how in r["decoys"]:
            decoys[cat][how] += 1
        for k in r["leaks"]:
            leaks[k] += 1
    return {"score": evaluate.score(predicted, gold), "outcomes": outcomes, "decoys": decoys, "leaks": leaks,
            "files": len(manifest["files"])}


# ---- report ---------------------------------------------------------------------------------------

def _fmt(x):
    return "  -  " if x is None else f"{x:5.2f}"


def format_report(res: dict, seconds: float) -> str:
    lines = [f"{res['files']} files, {seconds:.0f}s", "", "Per type (REVIEW findings are never true positives)",
            f"{'type':<16}{'gold':>5}{'TP':>5}{'FP':>5}{'FN':>5}{'review':>8}{'caught':>8}{'prec':>7}{'rec':>7}{'F1':>7}"]
    per = res["score"]["per_type"]
    for t in [x.value for x in SUPPORTED] + ["person", "address", "date_of_birth", "other"]:
        r = per.get(t)
        if r is None:
            continue
        lines.append(f"{t:<16}{r['tp'] + r['fn']:>5}{r['tp']:>5}{r['fp']:>5}{r['fn']:>5}{r['review']:>8}"
                     f"{r['review_caught']:>8}{_fmt(r['precision']):>7}{_fmt(r['recall']):>7}{_fmt(r['f1']):>7}")
    o = res["score"]["overall"]
    lines.append(f"{'overall':<16}{o['tp'] + o['fn']:>5}{o['tp']:>5}{o['fp']:>5}{o['fn']:>5}{o['review']:>8}"
                 f"{o['review_caught']:>8}{_fmt(o['precision']):>7}{_fmt(o['recall']):>7}{_fmt(o['f1']):>7}")
    lines += ["", "Look-alike decoys (should be clean; 'review' blocks release, 'masked' is a false positive)",
              f"{'decoy':<20}{'clean':>7}{'review':>8}{'masked':>8}{'lost':>6}"]
    for cat in sorted(res["decoys"]):
        c = res["decoys"][cat]
        lines.append(f"{cat:<20}{c['clean']:>7}{c['review']:>8}{c['masked']:>8}{c['unlocated']:>6}")
    lines += ["", "Pipeline outcome per format", f"{'format':<14}{'released':>9}{'blocked':>9}{'error':>7}"]
    for kind in KINDS:
        if kind in res["outcomes"]:
            c = res["outcomes"][kind]
            lines.append(f"{kind:<14}{c['released']:>9}{c['blocked']:>9}{c['error']:>7}")
    sup = {k: v for k, v in res["leaks"].items() if k in {t.value for t in SUPPORTED}}
    unsup = {k: v for k, v in res["leaks"].items() if k not in sup}
    lines += ["", f"Released files still containing a SUPPORTED-type value (must be 0): {sum(sup.values())}"
              + (f"  {sup}" if sup else ""),
              f"Released files still containing an unsupported-type value (names/addresses/DOB, by design): "
              f"{sum(unsup.values())}  {unsup if unsup else ''}"]
    return "\n".join(lines)


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(prog="pii_pipeline.benchmark")
    ap.add_argument("--out", type=Path, default=Path("tmp/benchmark"), help="gitignored work folder")
    ap.add_argument("--seed", type=int, default=1)
    ap.add_argument("--per-kind", type=int, default=9)
    ap.add_argument("--kinds", default=",".join(KINDS))
    ap.add_argument("--no-generate", action="store_true", help="reuse the corpus already in --out")
    args = ap.parse_args(argv)
    kinds = tuple(k for k in args.kinds.split(",") if k)
    bad = [k for k in kinds if k not in KINDS]
    if bad:
        print(f"error: unknown kinds {bad}; choose from {KINDS}", file=sys.stderr)
        return 2
    t0 = time.perf_counter()
    if not args.no_generate:
        generate(args.out, args.seed, args.per_kind, kinds)
    res = run_corpus(args.out)
    print(format_report(res, time.perf_counter() - t0))
    return 0


if __name__ == "__main__":
    sys.exit(main())
