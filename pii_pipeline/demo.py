"""One-command demo: python -m pii_pipeline.demo [--out tmp/demo]

Generates six small synthetic documents (reserved/test values only), runs each through the real CLI
entry point (pii_pipeline.cli.main), and prints what happened. The console output is meant to be
projected, so it never contains an original PII value: findings are shown as type + location + length,
and "after" is the placeholder or removal that replaced them. Whether the values are really gone is
checked by searching the released file for the known values and printing only the count.

Cases (expected outcome):
 1 DOCX  PII in body, table, header, hyperlink            -> RELEASED
 2 PPTX  PII in slide, table, speaker notes, comment      -> RELEASED (comment removed)
 3 PDF   native text layer                                -> RELEASED (true text removal)
 4 PDF   scanned page rendered as an image (needs OCR)    -> RELEASED, or skipped without Tesseract
 5 DOCX  order number (7+ digit run, needs human review)  -> BLOCKED
 6 DOCX  contains a picture (cannot be inspected)         -> BLOCKED
"""
import argparse
import contextlib
import io
import json
import shutil
import sys
import time
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

from . import cli, extract, ingest, package, pdf_extract

# Reserved / test values only (example.com/.org, 555-01xx, test card, TEST-NET addresses, SSN specimen).
EMAIL, PHONE, CARD, SSN, IP = "demo.person@example.com", "415-555-0142", "4111 1111 1111 1111", "123-45-6789", "203.0.113.42"
LINK_EMAIL = "demo.link@example.org"
EMAIL2, PHONE2, IP2 = "slides.owner@example.com", "212-555-0188", "198.51.100.17"
COMMENT_EMAIL, COMMENT_AUTHOR = "comment.reviewer@example.org", "Demo Reviewer"
EMAIL3, PHONE3, SSN3, CARD3 = "pdf.person@example.com", "310-555-0166", "321-54-9876", "5555 5555 5555 4444"
EMAIL4 = "scan.person@example.com"
ORDER = "ORD-2024-482913"

P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
FIXED_DATE = (2020, 1, 1, 0, 0, 0)


@dataclass
class Case:
    title: str
    path: Path
    expect_release: bool
    values: list[str]
    skip_reason: str | None = None


@dataclass
class Result:
    title: str
    expected: str
    actual: str
    ok: bool
    lines: list[str] = field(default_factory=list)


# ---- fixtures ---------------------------------------------------------------------------------------

def _rewrite_zip(path: Path, edit=None) -> None:
    """Rewrite an OPC zip with fixed timestamps (and optional edits) so the demo files are reproducible."""
    with zipfile.ZipFile(path) as z:
        parts = {n: z.read(n) for n in z.namelist()}
    if edit:
        edit(parts)
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for n, b in parts.items():
            z.writestr(zipfile.ZipInfo(n, FIXED_DATE), b, zipfile.ZIP_DEFLATED)


def _build_docx_pii(path: Path) -> None:
    from docx import Document
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml.ns import nsdecls
    from lxml import etree
    d = Document()
    d.add_paragraph(f"Contact {EMAIL} from host {IP} for details.")
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).text = f"Phone {PHONE}"
    t.cell(0, 1).text = f"Card {CARD}"
    d.sections[0].header.paragraphs[0].text = f"Employee SSN {SSN}"
    p = d.add_paragraph("Questions? ")
    rid = d.part.relate_to(f"mailto:{LINK_EMAIL}", RT.HYPERLINK, is_external=True)
    p._p.append(etree.fromstring(
        f'<w:hyperlink {nsdecls("w", "r")} r:id="{rid}"><w:r><w:t>write to us</w:t></w:r></w:hyperlink>'))
    d.save(path)
    _rewrite_zip(path)


def _build_pptx_pii(path: Path) -> None:
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.shapes.add_textbox(Inches(1), Inches(1), Inches(7), Inches(1)).text_frame.text = f"Reach us at {EMAIL2}"
    s.shapes.add_table(1, 2, Inches(1), Inches(3), Inches(7), Inches(1)).table.cell(0, 0).text = f"Call {PHONE2}"
    s.notes_slide.notes_text_frame.text = f"Speaker note: server {IP2}"
    prs.save(path)

    def add_comment(parts):
        parts["ppt/commentAuthors.xml"] = (
            f'<?xml version="1.0" encoding="UTF-8"?><p:cmAuthorLst xmlns:p="{P_NS}"><p:cmAuthor id="0" '
            f'name="{COMMENT_AUTHOR}" initials="DR" lastIdx="1" clrIdx="0"/></p:cmAuthorLst>').encode()
        parts["ppt/comments/comment1.xml"] = (
            f'<?xml version="1.0" encoding="UTF-8"?><p:cmLst xmlns:p="{P_NS}"><p:cm authorId="0" '
            f'dt="2026-01-01T00:00:00.000" idx="1"><p:pos x="10" y="10"/><p:text>Please email {COMMENT_EMAIL}'
            f'</p:text></p:cm></p:cmLst>').encode()
        for name, old, new in (
                ("ppt/_rels/presentation.xml.rels", "</Relationships>",
                 f'<Relationship Id="rIdCA" Type="{REL}/commentAuthors" Target="commentAuthors.xml"/></Relationships>'),
                ("ppt/slides/_rels/slide1.xml.rels", "</Relationships>",
                 f'<Relationship Id="rIdCM" Type="{REL}/comments" Target="../comments/comment1.xml"/></Relationships>'),
                ("[Content_Types].xml", "</Types>",
                 '<Override PartName="/ppt/commentAuthors.xml" ContentType="application/vnd.openxmlformats-officedocument.'
                 'presentationml.commentAuthors+xml"/><Override PartName="/ppt/comments/comment1.xml" ContentType='
                 '"application/vnd.openxmlformats-officedocument.presentationml.comments+xml"/></Types>')):
            parts[name] = parts[name].decode("utf-8").replace(old, new, 1).encode("utf-8")
    _rewrite_zip(path, add_comment)


def _build_native_pdf(path: Path) -> None:
    import pymupdf
    d = pymupdf.open()
    p = d.new_page()
    for k, line in enumerate(["Customer record (synthetic)", f"Email: {EMAIL3}", f"Phone: {PHONE3}",
                              f"SSN: {SSN3}", f"Card: {CARD3}", "Thank you for your business."]):
        p.insert_text((72, 100 + 24 * k), line, fontsize=12)
    d.save(str(path))


def _build_scanned_pdf(path: Path) -> None:
    import pymupdf
    from .benchmark import _png
    d = pymupdf.open()
    p = d.new_page()
    p.insert_image(p.rect, stream=_png(["Scanned letter (synthetic)", f"Please write to {EMAIL4} today.",
                                        "Thank you for your time."]))
    d.save(str(path))


def _build_docx_order(path: Path) -> None:
    from docx import Document
    d = Document()
    d.add_paragraph(f"Order {ORDER} has shipped.")
    d.save(path)
    _rewrite_zip(path)


def _build_docx_picture(path: Path, png: Path) -> None:
    from docx import Document
    from PIL import Image
    Image.new("RGB", (60, 60), "white").save(png)
    d = Document()
    d.add_paragraph("A report with a logo.")
    d.add_picture(str(png))
    d.save(path)
    _rewrite_zip(path)


# ---- checks and printing ------------------------------------------------------------------------------

def _present_values(path: Path, values: list[str]) -> int:
    """How many of the known values can still be found anywhere in the released file."""
    if path.suffix == ".pdf":
        import pymupdf
        d = pymupdf.open(str(path))
        blob = [str(d.metadata)]
        for page in d:
            blob.append(str(page.get_text("rawdict")))
            blob.append(" ".join(str(link) for link in page.get_links()))
        for x in range(1, d.xref_length()):
            try:
                blob.append(d.xref_object(x, compressed=False))
                if d.xref_is_stream(x):
                    blob.append(d.xref_stream(x).decode("latin-1", "ignore"))
            except Exception:
                pass
        text = "\n".join(blob)
        d.close()
    else:
        with zipfile.ZipFile(path) as z:
            text = "\n".join(z.read(n).decode("utf-8", "ignore") for n in z.namelist())
    return sum(1 for v in values if v in text)


def _after(fmt: str, unit_id: str) -> str:
    if unit_id.startswith("pkg:") and package.should_remove(unit_id[4:].split("#")[0]):
        return "removed together with its part"
    if fmt == "pdf":
        return "pixels blanked (white box)" if unit_id.endswith("/ocr") else "text removed (black box)"
    return None


def _describe(fmt: str, rep: dict) -> list[str]:
    out = []
    for f in rep["findings"]:
        span = "whole unit" if f["start"] is None else f"chars {f['start']}-{f['end']}"
        where = f["unit_id"] + (" [relationship target, e.g. hyperlink]" if "_rels/" in f["unit_id"] else "")
        n = "?" if f["start"] is None else f["end"] - f["start"]
        label = f["type"].upper()
        after = _after(fmt, f["unit_id"]) or f'"[{label}]"'
        if f["action"] in ("review", "keep"):
            after = "left in place; needs human review" if f["action"] == "review" else "left in place"
        out.append(f"      {f['type']:<12} {where} ({span})  {f['action']:<6} via {f['detector']}")
        out.append(f"                   before <{label}, {n} chars>  ->  after {after}")
    return out


def _run_case(i: int, total: int, case: Case, out_dir: Path) -> Result:
    fmt = case.path.suffix.lstrip(".")
    expected = "RELEASED" if case.expect_release else "BLOCKED"
    lines = [f"[{i}/{total}] {case.title}", f"  input   : {case.path.name} ({fmt}), expected {expected}"]
    if case.skip_reason:
        lines.append(f"  result  : skipped: {case.skip_reason}")
        return Result(case.title, expected, "SKIPPED", True, lines)
    buf_out, buf_err = io.StringIO(), io.StringIO()
    with contextlib.redirect_stdout(buf_out), contextlib.redirect_stderr(buf_err):
        rc = cli.main([str(case.path), "--out", str(out_dir)])
    doc_id = ingest.load(case.path).doc_id[:12]
    rpath = out_dir / f"{doc_id}.report.json"
    if rc == 2 or not rpath.exists():
        lines.append(f"  result  : ERROR (exit {rc}): {buf_err.getvalue().strip()[:200]}")
        return Result(case.title, expected, "ERROR", False, lines)
    rep = json.loads(rpath.read_text(encoding="utf-8"))
    gate = rep["gate"]
    final = out_dir / f"{doc_id}.redacted.{fmt}"
    lines.append(f"  gate    : {'PASS' if gate['passed'] else 'FAIL'} (exit {rc})")
    lines.append("  reasons : " + ("; ".join(gate["reasons"]) if gate["reasons"] else "none"))
    if rep["extraction_warnings"]:
        lines.append("  warnings: " + "; ".join(rep["extraction_warnings"]))
    lines.append(f"  by type : {rep['counts_by_type'] or '{}'}")
    lines.append(f"  by action: {rep['counts_by_action'] or '{}'}")
    if fmt == "pdf":
        ocr = [u for u in extract.extract(ingest.load(case.path)).units if u.unit_id.endswith("/ocr")]
        for u in ocr:
            lines.append(f"  OCR     : {u.unit_id} method={u.method.value} mean word confidence={u.ocr_confidence}")
    if rep["findings"]:
        lines.append("  findings (type, location, what happened; values are never printed):")
        lines += _describe(fmt, rep)
    actual = "RELEASED" if rc == 0 else "BLOCKED"
    if rc == 0 and final.exists():
        left = _present_values(final, case.values)
        lines.append(f"  output  : {final.name}")
        if fmt == "pdf":
            def native_chars(p):
                return sum(len(u.text) for u in extract.extract(ingest.load(p)).units if "/" not in u.unit_id)
            lines.append(f"  text layer: {native_chars(case.path)} characters before, {native_chars(final)} after")
        lines.append(f"  check   : {left} of {len(case.values)} known original values found in the released file")
        if fmt in ("docx", "pptx"):
            with zipfile.ZipFile(case.path) as a, zipfile.ZipFile(final) as b:
                gone = sorted(set(a.namelist()) - set(b.namelist()))
            if gone:
                lines.append("  removed parts: " + ", ".join(gone))
        ok = left == 0
    else:
        lines.append("  output  : no file kept")
        ok = not list(out_dir.glob(f"{doc_id}.*.{fmt}"))
    ok = ok and actual == expected
    return Result(case.title, expected, actual, ok, lines)


def run(out: Path) -> list[Result]:
    out = Path(out)
    inputs, outputs = out / "inputs", out / "output"
    for d in (inputs, outputs):
        shutil.rmtree(d, ignore_errors=True)
        d.mkdir(parents=True, exist_ok=True)
    _build_docx_pii(inputs / "demo_1_docx_pii.docx")
    _build_pptx_pii(inputs / "demo_2_pptx_pii.pptx")
    _build_native_pdf(inputs / "demo_3_native.pdf")
    ocr_ok = pdf_extract.ocr_available()
    if ocr_ok:
        _build_scanned_pdf(inputs / "demo_4_scanned.pdf")
    else:
        (inputs / "demo_4_scanned.pdf").write_bytes(b"")
    _build_docx_order(inputs / "demo_5_order_number.docx")
    _build_docx_picture(inputs / "demo_6_picture.docx", inputs / "logo.png")
    cases = [
        Case("DOCX with PII in body, table, header and hyperlink", inputs / "demo_1_docx_pii.docx", True,
             [EMAIL, IP, PHONE, CARD, SSN, LINK_EMAIL]),
        Case("PPTX with PII in slide, table, speaker notes and a comment", inputs / "demo_2_pptx_pii.pptx", True,
             [EMAIL2, PHONE2, IP2, COMMENT_EMAIL, COMMENT_AUTHOR]),
        Case("Native-text PDF with PII (true text removal)", inputs / "demo_3_native.pdf", True,
             [EMAIL3, PHONE3, SSN3, CARD3]),
        Case("Scanned PDF (rendered image) with an email, read by OCR", inputs / "demo_4_scanned.pdf", True, [EMAIL4],
             None if ocr_ok else "Tesseract not found (add it to PATH or set TESSERACT_CMD)"),
        Case("DOCX with an order number (needs human review)", inputs / "demo_5_order_number.docx", False, [ORDER]),
        Case("DOCX containing a picture (cannot be inspected)", inputs / "demo_6_picture.docx", False, []),
    ]
    return [_run_case(i, len(cases), c, outputs) for i, c in enumerate(cases, 1)]


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(prog="pii_pipeline.demo")
    ap.add_argument("--out", type=Path, default=Path("tmp/demo"), help="gitignored work folder")
    args = ap.parse_args(argv)
    t0 = time.perf_counter()
    print("PII pipeline demo: synthetic documents only; no original value is printed.\n")
    results = run(args.out)
    for r in results:
        print("\n".join(r.lines))
        print()
    print("Summary")
    for r in results:
        print(f"  {'ok ' if r.ok else 'BAD'} {r.actual:<9} (expected {r.expected})  {r.title}")
    bad = [r for r in results if not r.ok]
    print(f"\n{'ALL CASES AS EXPECTED' if not bad else str(len(bad)) + ' CASE(S) NOT AS EXPECTED'}"
          f" in {time.perf_counter() - t0:.0f}s; files are in {args.out}")
    return 0 if not bad else 1


if __name__ == "__main__":
    sys.exit(main())
