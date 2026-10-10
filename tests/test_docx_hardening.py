"""Regression tests for hidden-package PII, metadata, adjacent phones, draft artifacts, source integrity.
All fixtures are synthetic and built in tmp_path."""
import io
import zipfile

import pytest
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml.ns import nsdecls
from lxml import etree

from pii_pipeline import cli, detect, extract, gate, ingest, redact
from pii_pipeline.models import Action, PIIType
from tests.test_docx_slice import make_docx

LINK_EMAIL = "link.leak@example.com"
FIELD_PHONE = "212-555-0147"
DOCVAR_EMAIL = "docvar.leak@example.com"
META = ["Zed Authorperson", "Quill Manager", "Acme Hiddenco", "Catego Ryname", "IDENT-7781-X",
        "Draftstat Us", "Rita Reviewer", "Custom Propvalue", "Peopleonly Name", "Subj Ectline",
        "Key Wordsline", "Comm Entsline", "Lastmod Person"]
PII = [LINK_EMAIL, FIELD_PHONE, DOCVAR_EMAIL]


def mutate_zip(path, edit):
    with zipfile.ZipFile(path) as z:
        parts = {n: z.read(n) for n in z.namelist()}
    edit(parts)
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for n, b in parts.items():
            z.writestr(n, b)


def sub(parts, name, old, new):
    text = parts[name].decode("utf-8")
    assert old in text, (name, old)
    parts[name] = text.replace(old, new, 1).encode("utf-8")


def build_leaky(path):
    d = Document()
    d.add_paragraph("Plain body text.")
    d.part.relate_to(f"mailto:{LINK_EMAIL}", RT.HYPERLINK, is_external=True)
    p = d.add_paragraph()
    p._p.append(etree.fromstring(
        f'<w:r {nsdecls("w")}><w:instrText xml:space="preserve"> HYPERLINK "tel:{FIELD_PHONE}" </w:instrText></w:r>'))
    p2 = d.add_paragraph()
    p2._p.append(etree.fromstring(
        f'<w:ins {nsdecls("w")} w:id="9" w:author="Rita Reviewer" w:date="2026-01-01T00:00:00Z">'
        f'<w:r><w:t>added</w:t></w:r></w:ins>'))
    cp = d.core_properties
    cp.author, cp.last_modified_by, cp.title = "Zed Authorperson", "Lastmod Person", "Zed Authorperson"
    cp.category, cp.identifier, cp.content_status = "Catego Ryname", "IDENT-7781-X", "Draftstat Us"
    cp.subject, cp.keywords, cp.comments = "Subj Ectline", "Key Wordsline", "Comm Entsline"
    d.save(path)

    def edit(parts):
        sub(parts, "word/settings.xml", "</w:settings>",
            f'<w:docVars><w:docVar w:name="owner" w:val="{DOCVAR_EMAIL}"/></w:docVars></w:settings>')
        sub(parts, "docProps/app.xml", "</Properties>",
            "<Manager>Quill Manager</Manager><Company>Acme Hiddenco</Company></Properties>")
        parts["docProps/custom.xml"] = (
            '<?xml version="1.0" encoding="UTF-8"?><Properties xmlns="http://schemas.openxmlformats.org/'
            'officeDocument/2006/custom-properties"><property name="k" pid="2"><lpwstr>Custom Propvalue'
            '</lpwstr></property></Properties>').encode()
        sub(parts, "_rels/.rels", "</Relationships>",
            '<Relationship Id="rIdC1" Type="http://schemas.openxmlformats.org/officeDocument/2006/'
            'relationships/custom-properties" Target="docProps/custom.xml"/></Relationships>')
        sub(parts, "[Content_Types].xml", "</Types>",
            '<Override PartName="/docProps/custom.xml" ContentType="application/vnd.openxmlformats-'
            'officedocument.custom-properties+xml"/></Types>')
        parts["word/people.xml"] = (
            f'<?xml version="1.0" encoding="UTF-8"?><w15:people xmlns:w15="http://schemas.microsoft.com/'
            f'office/word/2012/wordml" {nsdecls("w")}><w15:person w15:author="Peopleonly Name"/></w15:people>').encode()
        sub(parts, "word/_rels/document.xml.rels", "</Relationships>",
            '<Relationship Id="rIdP1" Type="http://schemas.microsoft.com/office/2011/relationships/people" '
            'Target="people.xml"/></Relationships>')
        sub(parts, "[Content_Types].xml", "</Types>",
            '<Override PartName="/word/people.xml" ContentType="application/vnd.openxmlformats-'
            'officedocument.wordprocessingml.people+xml"/></Types>')

    mutate_zip(path, edit)
    return path


def package_blob(path):
    with zipfile.ZipFile(path) as z:
        return z.namelist(), b"\n".join(z.read(n) for n in z.namelist()).decode("utf-8", "ignore")


@pytest.fixture
def leaky_run(tmp_path):
    src = build_leaky(tmp_path / "leaky.docx")
    out = tmp_path / "out"
    code = cli.main([str(src), "--out", str(out)])
    return src, out, code


def test_leaky_source_findings_cover_hidden_parts(tmp_path):
    doc = extract.extract(ingest.load(build_leaky(tmp_path / "l.docx")))
    ids = {f.unit_id.split("#")[0] for f in detect.detect(doc)}
    assert "pkg:word/_rels/document.xml.rels" in ids
    assert "para:1" in ids                         # instrText now belongs to its paragraph's unit
    assert "pkg:word/settings.xml" in ids          # docVars
    d = gate.decide(doc, detect.detect(doc), doc, detect.detect(doc))
    assert not d.passed


@pytest.mark.parametrize("secret", PII + META)
def test_no_secret_survives_in_released_package(leaky_run, secret):
    _, out, code = leaky_run
    released = list(out.glob("*.redacted.docx"))
    assert code == 0 and len(released) == 1
    names, blob = package_blob(released[0])
    assert secret not in blob
    assert not any(n.startswith("docProps/thumbnail") for n in names)
    assert "docProps/custom.xml" not in names and "word/people.xml" not in names
    assert "docVar" not in blob and "<Manager>Quill" not in blob


def test_reports_and_output_never_contain_secrets(leaky_run, capsys):
    _, out, _ = leaky_run
    rep = next(out.glob("*.report.json")).read_text()
    assert not any(s in rep for s in PII + META)


def test_metadata_scrubbed_in_saved_package_xml(leaky_run):
    _, out, _ = leaky_run
    with zipfile.ZipFile(next(out.glob("*.redacted.docx"))) as z:
        core = etree.fromstring(z.read("docProps/core.xml"))
        for child in core:
            if etree.QName(child).localname not in ("created", "modified"):
                assert not (child.text or "").strip() and len(child) == 0
        app = z.read("docProps/app.xml").decode()
        assert "Manager" not in app and "Company" not in app


@pytest.mark.parametrize("text", ["Phones: 415 555 0132 212 555 0199", "Phones: 415-555-0132 212-555-0199"])
def test_adjacent_phones_both_detected(tmp_path, text):
    d = Document()
    d.add_paragraph(text)
    d.save(tmp_path / "p.docx")
    doc = extract.extract(ingest.load(tmp_path / "p.docx"))
    fs = [f for f in detect.detect(doc) if f.pii_type is PIIType.PHONE]
    assert len(fs) == 2 and all(f.action is Action.MASK for f in fs)
    assert cli.main([str(tmp_path / "p.docx"), "--out", str(tmp_path / "o")]) == 0
    _, blob = package_blob(next((tmp_path / "o").glob("*.redacted.docx")))
    assert "0199" not in blob and "0132" not in blob


def test_orphan_phone_fragment_cannot_pass(tmp_path):
    d = Document()
    d.add_paragraph("Call 555 0199 now")       # 7-digit local number: ambiguous
    d.save(tmp_path / "f.docx")
    fs = detect.detect(extract.extract(ingest.load(tmp_path / "f.docx")))
    assert fs and all(f.action is Action.REVIEW for f in fs)
    assert cli.main([str(tmp_path / "f.docx"), "--out", str(tmp_path / "o")]) == 1


def review_doc(tmp_path):
    d = Document()
    d.add_paragraph("Card 4111 1111 1111 1112 and mail jane.tester@example.com")
    d.save(tmp_path / "r.docx")
    return tmp_path / "r.docx"


def test_failed_gate_leaves_no_docx(tmp_path, capsys):
    out = tmp_path / "o"
    assert cli.main([str(review_doc(tmp_path)), "--out", str(out), "--audit"]) == 1
    assert list(out.glob("*.docx")) == [] and list(out.glob("*.audit.json")) == []
    assert len(list(out.glob("*.report.json"))) == 1
    text = capsys.readouterr().out
    assert "redacted=" not in text and "FAIL" in text and "jane.tester" not in text


def test_passing_run_leaves_only_released_file(tmp_path):
    out = tmp_path / "o"
    assert cli.main([str(make_docx(tmp_path / "ok.docx")), "--out", str(out)]) == 0
    assert [p.name.endswith(".redacted.docx") for p in out.glob("*.docx")] == [True]


def test_interrupted_run_leaves_no_artifact(tmp_path, monkeypatch):
    def boom(*a, **k):
        raise RuntimeError("simulated")
    monkeypatch.setattr(cli.gate, "decide", boom)
    out = tmp_path / "o"
    assert cli.main([str(make_docx(tmp_path / "ok.docx")), "--out", str(out)]) == 2
    assert list(out.glob("*.docx")) == []


def test_stale_release_removed_when_rerun_fails(tmp_path, monkeypatch):
    src = make_docx(tmp_path / "ok.docx")
    out = tmp_path / "o"
    assert cli.main([str(src), "--out", str(out)]) == 0
    monkeypatch.setattr(cli.gate, "decide", lambda *a, **k: gate.GateDecision(False, ["x"]))
    assert cli.main([str(src), "--out", str(out)]) == 1
    assert list(out.glob("*.docx")) == []


def test_redact_rejects_changed_source(tmp_path):
    src = make_docx(tmp_path / "a.docx")
    doc = extract.extract(ingest.load(src))
    findings = detect.detect(doc)
    make_docx(tmp_path / "b.docx")
    d = Document(str(tmp_path / "b.docx"))
    d.add_paragraph("changed after extraction")
    d.save(src)
    with pytest.raises(ValueError):
        redact.redact(doc, findings, tmp_path / "o")
    with pytest.raises(ValueError):
        extract.extract(doc.source)


def test_os_errors_do_not_leak_paths(tmp_path, capsys):
    blocker = tmp_path / "SECRETDIR_blocker"
    blocker.write_text("x")
    code = cli.main([str(make_docx(tmp_path / "ok.docx")), "--out", str(blocker / "sub")])
    err = capsys.readouterr().err
    assert code == 2 and "SECRETDIR" not in err and str(tmp_path) not in err


def test_uninspectable_binary_part_marks_incomplete(tmp_path):
    p = make_docx(tmp_path / "b.docx")
    mutate_zip(p, lambda parts: parts.update({"word/embeddings/blob.bin": b"\x00\x01"}))
    doc = extract.extract(ingest.load(p))
    assert not doc.complete
    assert cli.main([str(p), "--out", str(tmp_path / "o")]) == 1
