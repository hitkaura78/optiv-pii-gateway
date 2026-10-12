"""Split-node text, element-level coverage, removed-part findings, stale artifacts."""
import zipfile

import pytest
from docx import Document
from docx.oxml.ns import nsdecls
from lxml import etree

from pii_pipeline import cli, detect, extract, gate, ingest
from pii_pipeline.models import Action, PIIType
from tests.test_docx_hardening import mutate_zip, package_blob, sub
from tests.test_docx_slice import make_docx

W = nsdecls("w")
SPLIT_EMAIL = "split.leak@example.com"
SPLIT_PHONE = "212-555-0147"
CX_EMAIL = "cx.leak@example.com"
GLOSS_EMAIL = "gloss.leak@example.com"
CUSTOM_PROP_EMAIL = "customprop.leak@example.com"
SECRETS = [SPLIT_EMAIL, "mple.com", "split.leak", SPLIT_PHONE, "555", "0147", CX_EMAIL, GLOSS_EMAIL, CUSTOM_PROP_EMAIL]


def run_el(tag, text):
    return etree.fromstring(f'<w:r {W}><w:{tag} xml:space="preserve">{text}</w:{tag}></w:r>')


def build(path):
    d = Document()
    d.add_paragraph("body")
    p = d.add_paragraph()                                   # email split across instrText nodes
    for part in (' HYPERLINK "mailto:split.leak@exa', 'mple.com" '):
        p._p.append(run_el("instrText", part))
    p2 = d.add_paragraph()                                  # phone split across w:t and instrText
    p2._p.append(run_el("t", "Call 212-5"))
    p2._p.append(run_el("instrText", "55-"))
    p2._p.append(run_el("t", "0147 now"))
    body = d.element.body                                   # paragraph inside block-level customXml
    body.insert(len(body) - 1, etree.fromstring(
        f'<w:customXml {W} w:element="x"><w:p><w:r><w:t>cx.leak@exa</w:t></w:r>'
        f'<w:r><w:t>mple.com</w:t></w:r></w:p></w:customXml>'))
    d.save(path)

    def edit(parts):
        parts["word/glossary/document.xml"] = (
            f'<?xml version="1.0" encoding="UTF-8"?><w:glossaryDocument {W}><w:docParts><w:docPart><w:docPartBody>'
            f'<w:p><w:r><w:t>gloss.leak@</w:t></w:r><w:r><w:t>example.com</w:t></w:r></w:p>'
            f'</w:docPartBody></w:docPart></w:docParts></w:glossaryDocument>').encode()
    mutate_zip(path, edit)
    return path


def test_scan_sees_split_and_nested_content(tmp_path):
    doc = extract.extract(ingest.load(build(tmp_path / "s.docx")))
    found = {(f.pii_type, f.unit_id) for f in detect.detect(doc)}
    types = [t for t, _ in found]
    assert types.count(PIIType.EMAIL) >= 3 and PIIType.PHONE in types
    assert any(u.startswith("pkg:word/glossary/document.xml") for _, u in found)
    d = gate.decide(doc, detect.detect(doc), doc, detect.detect(doc))
    assert not d.passed


def test_split_and_nested_pii_redacted_from_every_fragment(tmp_path, capsys):
    out = tmp_path / "o"
    code = cli.main([str(build(tmp_path / "s.docx")), "--out", str(out)])
    assert code == 0
    released = list(out.glob("*.redacted.docx"))
    assert len(released) == 1
    with zipfile.ZipFile(released[0]) as z:
        assert z.testzip() is None
    _, blob = package_blob(released[0])
    for s in SECRETS:
        assert s not in blob, s
    again = extract.extract(ingest.load(released[0]))          # reopens and re-extracts
    assert not [f for f in detect.detect(again) if f.action is not Action.KEEP]
    shown = capsys.readouterr()
    text = shown.out + shown.err + next(out.glob("*.report.json")).read_text()
    assert not any(s in text for s in SECRETS if len(s) > 6 and s not in ("0147", "555"))   # hex ids may contain "555"


def test_unmappable_text_structure_fails_closed(tmp_path):
    d = Document()
    d.add_paragraph("x")
    d.element.body.insert(0, etree.fromstring(
        f'<w:p {W}><w:r><w:t>a@b.example<w:br/>.com</w:t></w:r></w:p>'))
    d.save(tmp_path / "u.docx")
    doc = extract.extract(ingest.load(tmp_path / "u.docx"))
    assert not doc.complete
    assert cli.main([str(tmp_path / "u.docx"), "--out", str(tmp_path / "o")]) == 1


def build_with_custom_prop(path):
    make_docx(path)

    def edit(parts):
        parts["docProps/custom.xml"] = (
            '<?xml version="1.0" encoding="UTF-8"?><Properties xmlns="http://schemas.openxmlformats.org/'
            f'officeDocument/2006/custom-properties"><property name="k" pid="2"><lpwstr>{CUSTOM_PROP_EMAIL}'
            '</lpwstr></property></Properties>').encode()
        sub(parts, "_rels/.rels", "</Relationships>",
            '<Relationship Id="rIdC1" Type="http://schemas.openxmlformats.org/officeDocument/2006/'
            'relationships/custom-properties" Target="docProps/custom.xml"/></Relationships>')
        sub(parts, "[Content_Types].xml", "</Types>",
            '<Override PartName="/docProps/custom.xml" ContentType="application/vnd.openxmlformats-'
            'officedocument.custom-properties+xml"/></Types>')
    mutate_zip(path, edit)
    return path


def test_finding_in_removed_part_is_accounted_for(tmp_path):
    src = build_with_custom_prop(tmp_path / "c.docx")
    doc = extract.extract(ingest.load(src))
    fs = [f for f in detect.detect(doc) if f.unit_id.startswith("pkg:docProps/custom.xml")]
    assert fs
    assert cli.main([str(src), "--out", str(tmp_path / "o")]) == 0
    out = next((tmp_path / "o").glob("*.redacted.docx"))
    names, blob = package_blob(out)
    assert CUSTOM_PROP_EMAIL not in blob and "docProps/custom.xml" not in names
    assert "custom.xml" not in blob          # relationship and content-type override gone
    with zipfile.ZipFile(out) as z:
        assert z.testzip() is None
    assert extract.extract(ingest.load(out)).complete


def test_stale_artifacts_for_same_doc_are_removed(tmp_path):
    src = make_docx(tmp_path / "ok.docx")
    stem = ingest.load(src).doc_id[:12]
    out = tmp_path / "o"
    out.mkdir()
    for n in (f"{stem}.redacted.docx", f"{stem}.UNSAFE-DRAFT.docx", f"{stem}.audit.json"):
        (out / n).write_text("stale")
    unrelated = out / "notes.txt"
    unrelated.write_text("keep")
    assert cli.main([str(src), "--out", str(out)]) == 0
    assert not (out / f"{stem}.audit.json").exists() and not (out / f"{stem}.UNSAFE-DRAFT.docx").exists()
    assert (out / f"{stem}.redacted.docx").read_bytes() != b"stale"
    assert unrelated.read_text() == "keep"


def test_stale_audit_removed_even_if_run_fails_early(tmp_path, monkeypatch):
    src = make_docx(tmp_path / "ok.docx")
    stem = ingest.load(src).doc_id[:12]
    out = tmp_path / "o"
    out.mkdir()
    (out / f"{stem}.audit.json").write_text("stale-sensitive")
    monkeypatch.setattr(cli.extract, "extract", lambda s: (_ for _ in ()).throw(ValueError("boom")))
    assert cli.main([str(src), "--out", str(out)]) == 2
    assert not (out / f"{stem}.audit.json").exists()


def test_gitignore_covers_drafts():
    text = open(".gitignore", encoding="utf-8").read()
    assert "*.UNSAFE-DRAFT*" in text


# ---- blocker A: sibling text in data XML; blocker B: prolog/epilog content and DOCTYPE ----
SIB_PARTS = ("<root><user>jane.te</user><dom>ster@example.com</dom>"
             "<a>415-555-</a><b>0132</b></root>")
PROLOG_EMAIL = "owner.jane.tester@example.com"


def add_part(path, name, xml):
    mutate_zip(path, lambda parts: parts.update({name: xml.encode("utf-8")}))
    return path


def released(out):
    return list(out.glob("*.docx"))


def test_sibling_split_custom_xml_blocks_release(tmp_path, capsys):
    src = add_part(make_docx(tmp_path / "s.docx"), "customXml/item1.xml", SIB_PARTS)
    doc = extract.extract(ingest.load(src))
    assert not doc.complete and doc.warnings
    out = tmp_path / "o"
    assert cli.main([str(src), "--out", str(out)]) == 1
    assert released(out) == []
    shown = capsys.readouterr()
    rep = next(out.glob("*.report.json")).read_text()
    for frag in ("jane.te", "ster@example.com", "415-555-"):   # short numerics could occur in hex ids
        assert frag not in shown.out + shown.err + rep


def test_single_field_custom_xml_still_redacted(tmp_path):
    src = add_part(make_docx(tmp_path / "s.docx"), "customXml/item1.xml",
                   "<root><user>solo.leak@example.com</user></root>")
    assert cli.main([str(src), "--out", str(tmp_path / "o")]) == 0
    assert "solo.leak" not in package_blob(released(tmp_path / "o")[0])[1]


@pytest.mark.parametrize("where", ["prolog", "epilog"])
def test_comment_outside_root_is_not_copied(tmp_path, where):
    p = make_docx(tmp_path / "c.docx")
    c = f"<!-- owner {PROLOG_EMAIL} -->"

    def edit(parts):
        text = parts["word/fontTable.xml"].decode("utf-8")
        i = text.index("<w:fonts")
        parts["word/fontTable.xml"] = ((text[:i] + c + text[i:]) if where == "prolog" else (text.rstrip() + c)).encode()
    mutate_zip(p, edit)
    doc = extract.extract(ingest.load(p))
    assert any(f.pii_type is PIIType.EMAIL and "fontTable" in f.unit_id for f in detect.detect(doc))
    assert cli.main([str(p), "--out", str(tmp_path / "o")]) == 0
    out = released(tmp_path / "o")[0]
    assert PROLOG_EMAIL not in package_blob(out)[1]
    with zipfile.ZipFile(out) as z:
        assert z.testzip() is None
    assert extract.extract(ingest.load(out)).complete


def test_doctype_rejected_and_no_artifact(tmp_path):
    p = make_docx(tmp_path / "d.docx")

    def edit(parts):
        text = parts["word/fontTable.xml"].decode("utf-8")
        i = text.index("<w:fonts")
        parts["word/fontTable.xml"] = (text[:i] + '<!DOCTYPE w:fonts [<!ENTITY x "y">]>' + text[i:]).encode()
    mutate_zip(p, edit)
    out = tmp_path / "o"
    assert cli.main([str(p), "--out", str(out)]) == 2
    assert released(out) == []


def test_every_xml_part_in_release_was_reserialized(tmp_path):
    """No unchecked original XML bytes: a PI outside the root of an untouched part is gone."""
    p = make_docx(tmp_path / "x.docx")
    mutate_zip(p, lambda parts: parts.update(
        {"word/webSettings.xml": b'<?xml version="1.0"?><?pi secret-marker?><w:webSettings '
                                 + nsdecls("w").encode() + b"/>"}))
    assert cli.main([str(p), "--out", str(tmp_path / "o")]) == 0
    assert "secret-marker" not in package_blob(released(tmp_path / "o")[0])[1]


# ---- metadata values and attribute-split fields ----
APP_NS = "http://schemas.openxmlformats.org/officeDocument/2006/extended-properties"
CORE_NS = ('xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties" '
           'xmlns:dc="http://purl.org/dc/elements/1.1/" xmlns:dcterms="http://purl.org/dc/terms/" '
           'xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance"')
GOOD_DATE = '<dcterms:created xsi:type="dcterms:W3CDTF">2026-01-01T00:00:00Z</dcterms:created>'
META_SECRETS = ["Jane Q Public", "Owner Name", "Loose Apptext", "Loose Coretext", "tail Coretail", "Attr Secret"]


def set_meta(path, app=None, core=None):
    def edit(parts):
        if app is not None:
            parts["docProps/app.xml"] = f'<?xml version="1.0"?><Properties xmlns="{APP_NS}">{app}</Properties>'.encode()
        if core is not None:
            parts["docProps/core.xml"] = f'<?xml version="1.0"?><cp:coreProperties {CORE_NS}>{core}</cp:coreProperties>'.encode()
    mutate_zip(path, edit)
    return path


@pytest.mark.parametrize("app,core", [
    ("<Pages>Jane Q Public</Pages>", None),
    ("<Words>Owner Name</Words>", None),
    ("Loose Apptext<Pages>3</Pages>tail Coretail", None),
    ('<Pages Note="Attr Secret">3</Pages>', None),
    (None, '<dcterms:created xsi:type="dcterms:W3CDTF">Jane Q Public</dcterms:created>'),
    (None, "Loose Coretext" + GOOD_DATE + "tail Coretail"),
    (None, '<dcterms:modified xsi:type="dcterms:W3CDTF" Note="Attr Secret">2026-01-01T00:00:00Z</dcterms:modified>'),
])
def test_metadata_values_are_validated_not_trusted(tmp_path, app, core):
    p = set_meta(make_docx(tmp_path / "m.docx"), app, core)
    assert cli.main([str(p), "--out", str(tmp_path / "o")]) == 0
    out = released(tmp_path / "o")[0]
    names, blob = package_blob(out)
    for sec in META_SECRETS:
        assert sec not in blob, sec
    with zipfile.ZipFile(out) as z:
        assert z.testzip() is None
    assert extract.extract(ingest.load(out)).complete
    Document(str(out))                                     # python-docx reopens it


def test_valid_statistics_and_dates_are_kept(tmp_path):
    p = set_meta(make_docx(tmp_path / "m.docx"), "<Pages>3</Pages><ScaleCrop>false</ScaleCrop>", GOOD_DATE)
    assert cli.main([str(p), "--out", str(tmp_path / "o")]) == 0
    blob = package_blob(released(tmp_path / "o")[0])[1]
    assert "<Pages>3</Pages>" in blob and "2026-01-01T00:00:00Z" in blob


def test_verify_scrubbed_checks_saved_bytes():
    from pii_pipeline import package
    bad = {"docProps/app.xml": f'<Properties xmlns="{APP_NS}"><Pages>Jane Q Public</Pages></Properties>'.encode(),
           "docProps/core.xml": f'<cp:coreProperties {CORE_NS}>{GOOD_DATE}</cp:coreProperties>'.encode()}
    with pytest.raises(ValueError):
        package.verify_scrubbed(list(bad), bad.__getitem__)
    bad["docProps/app.xml"] = f'<Properties xmlns="{APP_NS}"/>'.encode()
    bad["docProps/core.xml"] = f'<cp:coreProperties {CORE_NS}>loose<dcterms:created>x</dcterms:created></cp:coreProperties>'.encode()
    with pytest.raises(ValueError):
        package.verify_scrubbed(list(bad), bad.__getitem__)
    bad["docProps/core.xml"] = f'<cp:coreProperties {CORE_NS}>{GOOD_DATE}</cp:coreProperties>'.encode()
    package.verify_scrubbed(list(bad), bad.__getitem__)


def test_attribute_split_in_unknown_part_blocks_release(tmp_path):
    src = add_part(make_docx(tmp_path / "a.docx"), "customXml/item1.xml",
                   '<c><u a="jane.te" b="ster@example.com"/></c>')
    doc = extract.extract(ingest.load(src))
    assert not doc.complete
    out = tmp_path / "o"
    assert cli.main([str(src), "--out", str(out)]) == 1
    assert released(out) == []


def test_single_attribute_in_unknown_part_still_redacted(tmp_path):
    src = add_part(make_docx(tmp_path / "a.docx"), "customXml/item1.xml", '<c a="attr.leak@example.com"/>')
    assert cli.main([str(src), "--out", str(tmp_path / "o")]) == 0
    assert "attr.leak" not in package_blob(released(tmp_path / "o")[0])[1]


def test_known_word_parts_keep_independent_attributes(tmp_path):
    """Ordinary Word output (many attributes in styles/settings) must not be blocked."""
    assert extract.extract(ingest.load(make_docx(tmp_path / "k.docx"))).complete
