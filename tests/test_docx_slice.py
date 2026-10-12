import json
import zipfile

import pytest
from docx import Document
from docx.oxml.ns import nsdecls
from lxml import etree

from pii_pipeline import cli, detect, extract, gate, ingest
from pii_pipeline.models import Action, PIIType

EMAIL, PHONE, CARD, SSN, IP = "jane.tester@example.com", "415-555-0132", "4111 1111 1111 1111", "123-45-6789", "203.0.113.42"
ALL = [EMAIL, PHONE, CARD, SSN, IP, "jane.tester", "0132"]


def make_docx(path):
    d = Document()
    d.add_paragraph(f"Contact {PHONE} for details.")
    p = d.add_paragraph("Mail ")
    p.add_run("jane.tes").bold = True       # email split across runs
    p.add_run("ter@exam")
    p.add_run("ple.com today")
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).text = f"Card {CARD}"
    t.cell(0, 1).text = f"SSN {SSN} host {IP}"
    d.sections[0].header.paragraphs[0].text = f"Header {EMAIL}"
    d.save(path)
    return path


def raw_text(path):
    with zipfile.ZipFile(path) as z:
        return b"".join(z.read(n) for n in z.namelist()).decode("utf-8", "ignore")


def test_end_to_end_redaction_and_report(tmp_path):
    out = tmp_path / "out"
    decision, result, rpath, _ = cli.run(make_docx(tmp_path / "in.docx"), out)
    assert decision.passed, decision.reasons
    blob = raw_text(result.output_path)
    assert not any(s in blob for s in ALL)
    for tag in ("[EMAIL]", "[PHONE]", "[CREDIT_CARD]", "[US_SSN]", "[IP_ADDRESS]"):
        assert tag in blob
    rep = rpath.read_text()
    assert not any(s in rep for s in ALL)
    assert json.loads(rep)["counts_by_type"]["email"] == 2
    assert result.audit_path is None and not list(out.glob("*.audit.json"))


def test_extract_locations_and_header(tmp_path):
    doc = extract.extract(ingest.load(make_docx(tmp_path / "in.docx")))
    ids = [u.unit_id for u in doc.units]
    assert "para:1" in ids and "table:0/r0/c1/para:0" in ids and "section:0/header/para:0" in ids
    assert doc.complete


def test_split_run_email_detected(tmp_path):
    doc = extract.extract(ingest.load(make_docx(tmp_path / "in.docx")))
    f = [x for x in detect.detect(doc) if x.unit_id == "para:1"]
    assert [x.pii_type for x in f] == [PIIType.EMAIL]


def test_audit_is_opt_in(tmp_path):
    _, result, _, _ = cli.run(make_docx(tmp_path / "in.docx"), tmp_path / "o", write_audit=True)
    assert result.audit_path.name.endswith(".audit.json") and result.audit_path.exists()


def test_luhn_and_review(tmp_path):
    d = Document()
    d.add_paragraph("Card 4111 1111 1111 1112 and id 4155550132 here")
    d.save(tmp_path / "r.docx")
    # body only: style names elsewhere hit the known name over-masking (see test_m1_m3_integration)
    fs = [f for f in detect.detect(extract.extract(ingest.load(tmp_path / "r.docx")))
          if f.unit_id.startswith("para:")]
    assert fs and {f.action for f in fs} == {Action.REVIEW}
    decision, _, _, _ = cli.run(tmp_path / "r.docx", tmp_path / "o")
    assert not decision.passed and any("human review" in r for r in decision.reasons)
    assert cli.main([str(tmp_path / "r.docx"), "--out", str(tmp_path / "o")]) == 1


def test_gate_blocks_residual(tmp_path):
    doc = extract.extract(ingest.load(make_docx(tmp_path / "in.docx")))
    d = gate.decide(doc, detect.detect(doc), doc, detect.detect(doc))   # "redacted" still has PII
    assert not d.passed and any("residual" in r for r in d.reasons)


def test_unread_content_marks_incomplete(tmp_path):
    d = Document(str(make_docx(tmp_path / "in.docx")))
    d.element.body.insert(0, etree.fromstring(f"<w:p {nsdecls('w')}><w:r><w:drawing/></w:r></w:p>"))
    d.save(tmp_path / "img.docx")
    doc = extract.extract(ingest.load(tmp_path / "img.docx"))
    assert not doc.complete and doc.warnings
    assert cli.main([str(tmp_path / "img.docx"), "--out", str(tmp_path / "o")]) == 1


def test_ingest_rejections(tmp_path):
    (tmp_path / "e.docx").write_bytes(b"")
    (tmp_path / "t.docx").write_text("hello")
    (tmp_path / "z.docx").write_bytes(b"PK\x03\x04garbage")
    (tmp_path / "x.docx").write_bytes(b"%PDF-1.4\n")      # extension lies
    for name in ("e.docx", "t.docx", "z.docx", "x.docx", "missing.docx"):
        with pytest.raises(ValueError):
            ingest.load(tmp_path / name)
    ok = ingest.load(make_docx(tmp_path / "ok.docx"))
    assert len(ok.doc_id) == 64


def test_ingest_size_limit(tmp_path, monkeypatch):
    p = make_docx(tmp_path / "ok.docx")
    monkeypatch.setattr(ingest, "MAX_BYTES", 10)
    with pytest.raises(ValueError):
        ingest.load(p)
