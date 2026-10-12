"""Member 1 pipeline + Member 3 engine through the real CLI path. Synthetic values only."""
import zipfile

import pytest
from docx import Document

from pii_pipeline import cli, detect, extract, ingest
from pii_pipeline.models import Action, PIIType

# Hard cases from the guide: H3 prose names, H5 SSN area 900, H6 initials, H7 7 digits after +1,
# H8 last four, A4 employee/vendor IDs, DOB context
VALUES = ["Jane Doe", "EMP-40001", "14 July 1990", "900-55-2043", "+1 (555) 0114",
          "Rafael E. Duarte", "MER-IN-4471", "2218", "jane.doe@example.com"]


def _docx(path):
    d = Document()
    d.add_paragraph("Affected: Jane Doe (EMP-40001, DOB 14 July 1990), SSN 900-55-2043.")
    d.add_paragraph("Reach her on +1 (555) 0114 or jane.doe@example.com. Card ending 2218.")
    d.add_paragraph("Statement by Rafael E. Duarte, vendor MER-IN-4471, re INC-2026-0417.")
    d.save(path)
    return path


def _package_text(path):
    with zipfile.ZipFile(path) as z:
        return "".join(z.read(n).decode("utf-8", "replace") for n in z.namelist() if n.endswith(".xml"))


def test_cli_uses_member3_engine_and_releases_nothing_sensitive(tmp_path):
    decision, result, rpath, _ = cli.run(_docx(tmp_path / "in.docx"), tmp_path / "out")
    assert decision.passed, decision.reasons
    blob = _package_text(result.output_path)
    for v in VALUES:
        assert v not in blob, v
    assert "INC-2026-0417" in blob          # H9: allow-listed reference ID is kept, not masked
    report = rpath.read_text(encoding="utf-8")
    for v in VALUES:
        if len(v) >= 6:                       # "2218" can occur by chance inside hex ids
            assert v not in report, v
    assert "ALLOW_REF_ID_INC" in report       # every suppression is visible with its reason code


def test_safety_net_respects_engine_allowlist_and_keeps_unclaimed_runs(tmp_path):
    d = Document()
    d.add_paragraph("Ticket INC-2026-0417; reference 48291377 unexplained.")
    d.save(tmp_path / "n.docx")
    fs = detect.detect(extract.extract(ingest.load(tmp_path / "n.docx")))
    body = [f for f in fs if f.unit_id == "para:0"]
    assert any(f.action is Action.KEEP and getattr(f, "reason", "") == "ALLOW_REF_ID_INC" for f in body)
    # the INC digits are not sent to review; the unexplained run still is (fail closed)
    reviews = [f for f in body if f.action is Action.REVIEW]
    assert len(reviews) == 1 and reviews[0].pii_type is PIIType.OTHER


@pytest.mark.xfail(strict=True, reason="KNOWN LIMITATION: without spaCy the name heuristic treats any "
                   "two capitalised words as a person, e.g. Word style names; fails safe (over-masks)")
def test_style_names_are_not_masked_as_people(tmp_path):
    Document().save(tmp_path / "e.docx")
    fs = detect.detect(extract.extract(ingest.load(tmp_path / "e.docx")))
    assert not [f for f in fs if f.pii_type is PIIType.PERSON and f.action is not Action.KEEP]


def test_residual_scan_catches_value_left_in_redacted_output(tmp_path):
    d = Document()
    d.add_paragraph("Left over: leaked.user@example.org")
    d.save(tmp_path / "r.docx")
    residual = detect.detect_residual(extract.extract(ingest.load(tmp_path / "r.docx")))
    assert any(f.pii_type is PIIType.EMAIL and f.action is not Action.KEEP for f in residual)
