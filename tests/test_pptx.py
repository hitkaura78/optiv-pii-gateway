"""PPTX vertical slice: every text location, metadata, fail-closed parts, CLI behaviour.
Synthetic fixtures built with python-pptx in tmp_path only."""
import io
import zipfile

import pytest
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

from pii_pipeline import cli, detect, extract, ingest
from pii_pipeline.models import PIIType
from tests.test_docx_hardening import mutate_zip, package_blob, sub

EMAIL, PHONE = "jane.tester@example.com", "415-555-0132"
LEAKS = [EMAIL, "jane.tes", "ter@exam", PHONE, "0132", "grp.leak", "notes.leak", "hidden.leak", "layout.leak",
         "link.leak", "master.leak", "alt.leak", "comment.leak"]
META = ["Zed Authorperson", "Lastmod Person", "Catego Ryname", "Quill Manager", "Acme Hiddenco",
        "Custom Propvalue", "Pat Reviewer", "Subj Ectline"]
P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def make_pptx(path, pii=True):
    prs = Presentation()
    blank = prs.slide_layouts[6]
    s = prs.slides.add_slide(blank)
    if pii:
        tb = s.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
        p = tb.text_frame.paragraphs[0]
        for t in ("Mail jane.tes", "ter@exam", "ple.com now"):          # split across runs
            p.add_run().text = t
        gf = s.shapes.add_table(1, 2, Inches(1), Inches(3), Inches(6), Inches(1))
        gf.table.cell(0, 0).text = f"Call {PHONE}"
        grp = s.shapes.add_group_shape()
        grp.shapes.add_textbox(Inches(1), Inches(5), Inches(4), Inches(1)).text_frame.text = "grp.leak@example.com"
        s.notes_slide.notes_text_frame.text = "notes.leak@example.com"
        link = s.shapes.add_textbox(Inches(1), Inches(6), Inches(4), Inches(1))
        run = link.text_frame.paragraphs[0].add_run()
        run.text = "contact"
        run.hyperlink.address = "mailto:link.leak@example.com"
        alt = s.shapes.add_textbox(Inches(1), Inches(7), Inches(4), Inches(0.5))
        alt._element.nvSpPr.cNvPr.set("descr", "alt.leak@example.com")
        alt.text_frame.text = "alt"
        hidden = prs.slides.add_slide(blank)
        hidden._element.set("show", "0")
        hidden.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text_frame.text = "hidden.leak@example.com"
        prs.slide_layouts[0].placeholders[0].text_frame.text = "layout.leak@example.com"
        prs.slide_master.placeholders[0].text_frame.text = "master.leak@example.com"
    cp = prs.core_properties
    cp.author, cp.last_modified_by, cp.category, cp.subject = "Zed Authorperson", "Lastmod Person", "Catego Ryname", "Subj Ectline"
    prs.save(path)
    return path


def add_comment_parts(path):
    def edit(parts):
        parts["ppt/commentAuthors.xml"] = (
            f'<?xml version="1.0" encoding="UTF-8"?><p:cmAuthorLst xmlns:p="{P_NS}">'
            f'<p:cmAuthor id="0" name="Pat Reviewer" initials="PR" lastIdx="1" clrIdx="0"/></p:cmAuthorLst>').encode()
        parts["ppt/comments/comment1.xml"] = (
            f'<?xml version="1.0" encoding="UTF-8"?><p:cmLst xmlns:p="{P_NS}"><p:cm authorId="0" '
            f'dt="2026-01-01T00:00:00.000" idx="1"><p:pos x="10" y="10"/><p:text>comment.leak@example.com'
            f'</p:text></p:cm></p:cmLst>').encode()
        sub(parts, "ppt/_rels/presentation.xml.rels", "</Relationships>",
            f'<Relationship Id="rIdCA" Type="{REL}/commentAuthors" Target="commentAuthors.xml"/></Relationships>')
        sub(parts, "ppt/slides/_rels/slide1.xml.rels", "</Relationships>",
            f'<Relationship Id="rIdCM" Type="{REL}/comments" Target="../comments/comment1.xml"/></Relationships>')
        sub(parts, "[Content_Types].xml", "</Types>",
            '<Override PartName="/ppt/commentAuthors.xml" ContentType="application/vnd.openxmlformats-officedocument.'
            'presentationml.commentAuthors+xml"/><Override PartName="/ppt/comments/comment1.xml" ContentType='
            '"application/vnd.openxmlformats-officedocument.presentationml.comments+xml"/></Types>')
    mutate_zip(path, edit)


def add_part(path, name, data=b"x"):
    mutate_zip(path, lambda parts: parts.update({name: data}))
    return path


def pptx_files(out):
    return list(out.glob("*.pptx"))


def blob_of(path):
    return package_blob(path)[1]


def test_units_have_readable_locators(tmp_path):
    doc = extract.extract(ingest.load(make_pptx(tmp_path / "a.pptx")))
    ids = {u.unit_id for u in doc.units}
    assert "slide:1/shape:0/para:0" in ids and "slide:1/notes/para:0" in ids
    assert any(i.startswith("slide:1/shape:1/table/r0/c0/para:") for i in ids)
    assert doc.complete, doc.warnings
    types = [f.pii_type for f in detect.detect(doc)]
    assert types.count(PIIType.EMAIL) >= 8 and PIIType.PHONE in types


def test_every_text_location_is_redacted(tmp_path, capsys):
    src = make_pptx(tmp_path / "a.pptx")
    out = tmp_path / "o"
    assert cli.main([str(src), "--out", str(out)]) == 0
    rel = pptx_files(out)
    assert len(rel) == 1 and rel[0].name.endswith(".redacted.pptx")
    blob = blob_of(rel[0])
    for s in LEAKS:
        assert s not in blob, s
    with zipfile.ZipFile(rel[0]) as z:
        assert z.testzip() is None
    prs = Presentation(str(rel[0]))                       # python-pptx reopens it
    assert len(prs.slides) == 2
    again = extract.extract(ingest.load(rel[0]))
    assert again.complete and not detect.detect(again)
    shown = capsys.readouterr()
    rep = next(out.glob("*.report.json")).read_text()
    assert not any(s in shown.out + shown.err + rep for s in (EMAIL, PHONE, "grp.leak", "notes.leak"))
    assert not list(out.glob("*UNSAFE-DRAFT*")) and not list(out.glob("*.audit.json"))


def test_comments_and_authors_are_removed(tmp_path):
    src = make_pptx(tmp_path / "a.pptx")
    add_comment_parts(src)
    doc = extract.extract(ingest.load(src))
    assert any(u.unit_id.startswith("pkg:ppt/comments/") for u in doc.units)       # scanned before removal
    assert cli.main([str(src), "--out", str(tmp_path / "o")]) == 0
    out = pptx_files(tmp_path / "o")[0]
    names, blob = package_blob(out)
    assert "ppt/commentAuthors.xml" not in names and not any(n.startswith("ppt/comments/") for n in names)
    assert "Pat Reviewer" not in blob and "comment.leak" not in blob
    assert "commentAuthors" not in blob and "comments/comment1" not in blob       # rels + overrides gone
    Presentation(str(out))


def test_metadata_is_scrubbed_in_saved_package(tmp_path):
    src = make_pptx(tmp_path / "a.pptx", pii=False)

    def edit(parts):
        sub(parts, "docProps/app.xml", "</Properties>", "<Manager>Quill Manager</Manager><Company>Acme Hiddenco</Company></Properties>")
        parts["docProps/custom.xml"] = (
            '<?xml version="1.0"?><Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/'
            'custom-properties"><property name="k" pid="2"><lpwstr>Custom Propvalue</lpwstr></property></Properties>').encode()
        sub(parts, "_rels/.rels", "</Relationships>",
            f'<Relationship Id="rIdC1" Type="{REL}/custom-properties" Target="docProps/custom.xml"/></Relationships>')
        sub(parts, "[Content_Types].xml", "</Types>",
            '<Override PartName="/docProps/custom.xml" ContentType="application/vnd.openxmlformats-officedocument.'
            'custom-properties+xml"/></Types>')
    mutate_zip(src, edit)
    assert cli.main([str(src), "--out", str(tmp_path / "o")]) == 0
    out = pptx_files(tmp_path / "o")[0]
    names, blob = package_blob(out)
    for s in META:
        assert s not in blob, s
    assert "docProps/custom.xml" not in names and not any(n.startswith("docProps/thumbnail") for n in names)
    Presentation(str(out))


def test_plain_deck_without_pii_passes(tmp_path):
    src = make_pptx(tmp_path / "plain.pptx", pii=False)
    assert extract.extract(ingest.load(src)).complete
    assert cli.main([str(src), "--out", str(tmp_path / "o")]) == 0


def test_picture_blocks_release(tmp_path):
    img = io.BytesIO()
    Image.new("RGB", (40, 40), "white").save(img, "PNG")
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.shapes.add_picture(io.BytesIO(img.getvalue()), Inches(1), Inches(1))
    prs.save(tmp_path / "pic.pptx")
    assert not extract.extract(ingest.load(tmp_path / "pic.pptx")).complete
    assert cli.main([str(tmp_path / "pic.pptx"), "--out", str(tmp_path / "o")]) == 1
    assert pptx_files(tmp_path / "o") == []


def test_chart_blocks_release(tmp_path):
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    data = CategoryChartData()
    data.categories = ["a", "b"]
    data.add_series("S", (1, 2))
    s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1), Inches(4), Inches(3), data)
    prs.save(tmp_path / "chart.pptx")
    doc = extract.extract(ingest.load(tmp_path / "chart.pptx"))
    assert not doc.complete and any("chart" in w.lower() for w in doc.warnings)
    assert cli.main([str(tmp_path / "chart.pptx"), "--out", str(tmp_path / "o")]) == 1
    assert pptx_files(tmp_path / "o") == []


@pytest.mark.parametrize("name,data", [
    ("ppt/diagrams/data1.xml", b'<?xml version="1.0"?><d/>'),
    ("ppt/embeddings/oleObject1.bin", b"\x00\x01"),
    ("ppt/activeX/activeX1.xml", b'<?xml version="1.0"?><ax/>'),
    ("ppt/vbaProject.bin", b"\x00\x01"),
    ("ppt/media/audio1.wav", b"RIFF"),
], ids=["smartart", "embedded", "activex", "vba", "media"])
def test_unreadable_parts_block_release(tmp_path, name, data):
    src = add_part(make_pptx(tmp_path / "a.pptx", pii=False), name, data)
    assert not extract.extract(ingest.load(src)).complete
    assert cli.main([str(src), "--out", str(tmp_path / "o")]) == 1
    assert pptx_files(tmp_path / "o") == []


def test_macro_enabled_content_type_and_external_ole_link_block(tmp_path):
    src = make_pptx(tmp_path / "m.pptx", pii=False)
    def edit(parts):
        text = parts["[Content_Types].xml"].decode()
        parts["[Content_Types].xml"] = text.replace(
            "presentationml.presentation.main+xml", "ms-powerpoint.presentation.macroEnabled.main+xml").encode()
    mutate_zip(src, edit)
    out = tmp_path / "o"
    assert cli.main([str(src), "--out", str(out)]) in (1, 2)          # python-pptx itself refuses macro content
    assert pptx_files(out) == []
    src2 = make_pptx(tmp_path / "e.pptx", pii=False)
    mutate_zip(src2, lambda parts: sub(
        parts, "ppt/slides/_rels/slide1.xml.rels", "</Relationships>",
        f'<Relationship Id="rIdX" Type="{REL}/oleObject" Target="file:///C:/secret/x.xlsx" TargetMode="External"/>'
        f'</Relationships>'))
    assert not extract.extract(ingest.load(src2)).complete
    assert cli.main([str(src2), "--out", str(tmp_path / "o")]) == 1


def test_unknown_part_with_several_attributes_blocks(tmp_path):
    src = add_part(make_pptx(tmp_path / "a.pptx", pii=False), "ppt/tags/tag1.xml",
                   f'<p:tagLst xmlns:p="{P_NS}"><p:tag name="jane.te" val="ster@example.com"/></p:tagLst>'.encode())
    assert not extract.extract(ingest.load(src)).complete
    assert cli.main([str(src), "--out", str(tmp_path / "o")]) == 1


def test_review_finding_blocks_and_stale_pptx_artifacts_are_cleaned(tmp_path):
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text_frame.text = "Card 4111 1111 1111 1112"
    prs.save(tmp_path / "r.pptx")
    stem = ingest.load(tmp_path / "r.pptx").doc_id[:12]
    out = tmp_path / "o"
    out.mkdir()
    for n in (f"{stem}.redacted.pptx", f"{stem}.UNSAFE-DRAFT.pptx", f"{stem}.audit.json"):
        (out / n).write_text("stale")
    (out / "notes.txt").write_text("keep")
    assert cli.main([str(tmp_path / "r.pptx"), "--out", str(out), "--audit"]) == 1
    assert sorted(p.name for p in out.iterdir() if p.suffix != ".json") == ["notes.txt"]
    assert not list(out.glob("*.audit.json"))


def test_bad_inputs_exit_2(tmp_path):
    (tmp_path / "x.pptx").write_bytes(b"PK\x03\x04garbage")
    make_pptx(tmp_path / "real.pptx", pii=False)
    (tmp_path / "wrong.docx").write_bytes((tmp_path / "real.pptx").read_bytes())     # extension lies
    for name in ("x.pptx", "wrong.docx", "missing.pptx"):
        assert cli.main([str(tmp_path / name), "--out", str(tmp_path / "o")]) == 2
    assert pptx_files(tmp_path / "o") == []
